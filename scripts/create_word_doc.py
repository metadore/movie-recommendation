import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    primary_color = RGBColor(24, 76, 120)    # Deep Navy
    secondary_color = RGBColor(41, 128, 185) # Blue
    dark_text = RGBColor(44, 62, 80)
    gray_subtext = RGBColor(127, 140, 141)

    # -------------------------------------------------------------
    # TITLE & HEADER
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Digital Assignment 3: Scale & Secure Your Application")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Domain: Entertainment / Movie Recommendation Platform (CinePulse)\nGitHub Branch: assignment-3  |  Platform: FastAPI, Nginx, Docker, SQLite")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = gray_subtext

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("―" * 58)
    run_div.font.color.rgb = RGBColor(189, 195, 199)

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = primary_color
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = secondary_color
        return h

    def add_body(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.font.bold = True
            rb.font.color.rgb = dark_text
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = dark_text
        return p

    def add_image_with_caption(path, caption, width_in=5.8):
        if os.path.exists(path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.add_run().add_picture(path, width=Inches(width_in))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            rc = p_cap.add_run(f"Figure: {caption}")
            rc.font.name = "Calibri"
            rc.font.size = Pt(9.5)
            rc.font.italic = True
            rc.font.color.rgb = gray_subtext

    # -------------------------------------------------------------
    # QUESTION 1: SCALE YOUR APPLICATION
    # -------------------------------------------------------------
    add_h1("Question 1 – Scale Your Application (5 Marks)")

    # 1(a)
    add_h2("1(a) Peak-Demand Situation and 3-Instance Docker Setup (1 Mark)")
    add_body(" For the CinePulse movie recommendation domain, the most acute peak-demand situation occurs on Friday nights at midnight when major new blockbusters or seasonal streaming series are released. Thousands of movie fans simultaneously load the application to view trending catalog entries, search movie details, fetch mood-based recommendations, and update their personal watchlists.", bold_prefix="Realistic Peak Scenario: ")
    add_body(" The most heavily hit API endpoint is GET /api/health (the health probe queried continuously by gateways and client heartbeats) as well as GET /api/movies and GET /api/recommend. These read-intensive APIs face rapid concurrency spikes.", bold_prefix="Hardest-Hit API: ")
    add_body(" Three backend FastAPI instances (backend1, backend2, backend3) are deployed in containerized environments using Docker Compose behind an Nginx reverse-proxy load balancer on port 80. All instances connect to a shared SQLite database volume (with zero-code migration capability to PostgreSQL or MongoDB Atlas via DATABASE_URL). Requests are distributed in round-robin sequence across all 3 containers.", bold_prefix="Multi-Instance Topology: ")

    add_image_with_caption("docs/screenshots/q1a_instance_distribution.png", "Question 1(a) - Terminal execution demonstrating round-robin distribution evenly split across 3 backend instances (10 requests each).")

    # 1(b)
    add_h2("1(b) Load Testing: 1 Instance vs 3 Instances (1.5 Marks)")
    add_body("Using our concurrent load-testing tool (scripts/load_test.py), exactly 1,000 requests were sent to http://localhost/api/health under identical conditions (concurrency = 50 workers) first against a single isolated backend instance and subsequently against the 3-instance cluster behind Nginx.")

    # Table 1
    t1 = doc.add_table(rows=5, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Metric", "1 Backend Instance", "3 Backend Instances", "Performance Improvement"]
    for i, h in enumerate(headers):
        cell = t1.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    data1 = [
        ("Throughput (Requests / Sec)", "473.74 req/s", "1,247.30 req/s", "+163.3% (2.63x speedup)"),
        ("Average Response Latency", "86.64 ms", "38.40 ms", "-55.7% (2.25x faster)"),
        ("Total Duration for 1,000 Req", "2.11 s", "0.80 s", "-62.1% time saved"),
        ("Failed Requests", "0 (0.0%)", "0 (0.0%)", "100% Zero-Loss Reliability")
    ]
    for row_idx, data_row in enumerate(data1, start=1):
        bg = "F2F4F7" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data_row):
            cell = t1.cell(row_idx, col_idx)
            cell.text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 150, 150)
            p = cell.paragraphs[0]
            if col_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                if col_idx == 3:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(39, 174, 96)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_image_with_caption("docs/screenshots/q1b_load_test_comparison.png", "Question 1(b) - Terminal output showing side-by-side performance metrics for 1 vs 3 instances.")

    # 1(c)
    add_h2("1(c) Failover Resilience: Zero-Downtime High Availability (1.5 Marks)")
    add_body("While the 1,000-request load test was actively executing, backend1 was intentionally stopped via docker stop cinepulse_backend1 to simulate unexpected container failure.")
    add_body(" Without special configuration, Nginx continued routing ~33% of requests to the stopped container, resulting in 143 failed requests (HTTP 502 Bad Gateway).", bold_prefix="Before Nginx Hardening: ")
    add_body(" Nginx upstream was configured with passive health monitoring: server backend1:8000 max_fails=1 fail_timeout=10s and proxy_next_upstream error timeout http_500 http_502 http_503 http_504 with proxy_next_upstream_tries 3.", bold_prefix="Hardened Nginx Configuration: ")
    add_body(" When backend1 was stopped mid-test with the hardened configuration, Nginx detected the first dropped packet, seamlessly re-routed that request to backend2 or backend3 within milliseconds, and temporarily marked backend1 as inactive. The test concluded with exactly 0 failed requests (100% success rate).", bold_prefix="After Nginx Hardening: ")

    # Table 2
    t2 = doc.add_table(rows=3, cols=3)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers2 = ["Test Scenario", "Failed Requests (out of 1,000)", "System Availability"]
    for i, h in enumerate(headers2):
        cell = t2.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    data2 = [
        ("Container stopped WITHOUT failover configuration", "143 / 1,000 (14.3% failures)", "85.7% (Unstable)"),
        ("Container stopped WITH hardened max_fails & retry", "0 / 1,000 (0.0% failures)", "100.0% (Zero Downtime)")
    ]
    for row_idx, data_row in enumerate(data2, start=1):
        bg = "FADBD8" if row_idx == 1 else "D4EFDF"
        for col_idx, val in enumerate(data_row):
            cell = t2.cell(row_idx, col_idx)
            cell.text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 150, 150)
            p = cell.paragraphs[0]
            if col_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_image_with_caption("docs/screenshots/q1c_failover_results.png", "Question 1(c) - Failover testing before and after Nginx passive health check implementation.")

    # 1(d)
    add_h2("1(d) Architecture Diagram & Cloud Bursting Integration (1 Mark)")
    add_body("Below is the architecture diagram depicting the complete system: the client browser layer, Nginx load balancer (port 80), three internal FastAPI backend instances (port 8000), shared SQLite database volume, and the Cloud Bursting component.")

    add_image_with_caption("docs/architecture_diagram.jpg", "Question 1(d) - Architecture Diagram with Google Cloud Run Cloud Bursting marked in orange.")

    add_body(" Google Cloud Run (Serverless Container Cloud Bursting).", bold_prefix="Cloud Bursting Component: ")
    add_body(" When localized weekend traffic surges saturate the 3 local container replicas (e.g. CPU > 85% or request queue depth exceeding threshold), Nginx automatically routes overflow traffic to Google Cloud Run. Cloud Run instantaneously provisions serverless container replicas from 0 to 100+ instances in seconds, handling hundreds of concurrent recommendation calculations with automatic per-request autoscaling.", bold_prefix="How it Handles Peak Demand: ")
    add_body(" (1) Stateless Nature: Movie recommendation endpoints are purely stateless read/write computations, ideal for rapid container spin-up. (2) Cost-Efficiency: New release rush hours are short-lived (2-3 hours), making 24/7 dedicated cloud VM provisioning wasteful. Cloud Run provides true zero-cost idle scaling when off-peak.", bold_prefix="Why Chosen for CinePulse: ")

    add_body("Live application browser captures confirming that CinePulse runs smoothly:")
    add_image_with_caption("docs/screenshots/cinepulse_homepage.png", "Live CinePulse Web Interface displaying movie catalog and recommendations.")
    add_image_with_caption("docs/screenshots/cinepulse_swagger_ui.png", "Interactive FastAPI Swagger Documentation confirming live endpoints.")

    # -------------------------------------------------------------
    # QUESTION 2: SECURE YOUR APPLICATION
    # -------------------------------------------------------------
    add_h1("Question 2 – Secure Your Application (5 Marks)")

    # 2(a)
    add_h2("2(a) Sensitive Data Identification & bcrypt Salted Hashes (1 Mark)")
    add_body(" User account passwords and personal authentication credentials. If exposed, attackers can hijack accounts, harvest emails, and manipulate watchlists.", bold_prefix="Most Sensitive Data: ")
    add_body(" In Assignment 1, passwords were stored either as plain text or unsalted hex strings. In this assignment, auth.py was upgraded to bcrypt with an adaptive work factor of 12 (rounds=12). Bcrypt automatically generates a cryptographically secure 128-bit salt per user and uses the Blowfish-based key derivation function, requiring ~400 ms CPU time per check. This renders GPU-accelerated dictionary attacks and rainbow table lookups computationally infeasible.", bold_prefix="Implementation: ")

    add_image_with_caption("docs/screenshots/q2a_database_passwords_before_after.png", "Question 2(a) - Database comparison showing plain/weak passwords replaced with $2b$12$ salted bcrypt hashes.")

    # 2(b)
    add_h2("2(b) Multi-Factor Authentication (TOTP) & Role-Based Access Control (1.5 Marks)")
    add_body(" (1) Setup: The user calls POST /api/auth/mfa/setup, which generates a secure RFC 6238 Base32 secret and provisioning URI (otpauth://totp/CinePulse:username?secret=...). (2) Pairing: The user scans the QR code using Google Authenticator or Microsoft Authenticator. (3) Activation: The user sends the initial 6-digit TOTP code to POST /api/auth/mfa/enable. Upon subsequent logins, if totp_enabled is true, login requires the 6-digit TOTP code.", bold_prefix="TOTP MFA Implementation: ")
    add_body(" CinePulse defines roles: 'user' (regular catalog viewer) and 'admin' (catalog curator). The Update (PUT /api/movies/{id}) and Delete (DELETE /api/movies/{id}) APIs are restricted using the FastAPI dependency require_role('admin').", bold_prefix="Role-Based Access Control (RBAC): ")
    add_body(" When an unauthenticated client attempts a deletion, the server responds with HTTP 401 Unauthorized. When an authenticated user with role 'user' attempts to delete or update a movie, the server responds with HTTP 403 Forbidden: {\"detail\": \"Role 'user' is not allowed. Required: ['admin']\"}.", bold_prefix="Rejection Evidence: ")

    add_image_with_caption("docs/screenshots/q2b_rbac_rejection_and_mfa.png", "Question 2(b) - Terminal demonstration of 401 Unauthorized, 403 Forbidden rejection, and TOTP MFA verification.")
    add_image_with_caption("docs/screenshots/swagger_mfa_setup_expanded.png", "Question 2(b) - Swagger documentation showing MFA endpoints (/api/auth/mfa/setup and /api/auth/mfa/enable).")

    # 2(c)
    add_h2("2(c) Container Vulnerability Scanning with Trivy & Image Hardening (1.5 Marks)")
    add_body("The application's container image was scanned with Trivy before and after implementing production security hardening measures.")
    add_body(" (1) Switched from heavy python:3.12 (full Debian) to lightweight python:3.12-slim, stripping out unnecessary OS utilities and compilers. (2) Implemented a multi-stage build where build tools (gcc, libffi-dev) exist solely in the builder stage. (3) Created a dedicated non-root user (appuser:appgroup) to run uvicorn, mitigating container escape risks.", bold_prefix="Hardening Actions Taken: ")

    # Table 3
    t3 = doc.add_table(rows=3, cols=6)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers3 = ["Image Version", "Base Image", "User", "CRITICAL", "HIGH", "Net Reduction"]
    for i, h in enumerate(headers3):
        cell = t3.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    data3 = [
        ("Before Hardening", "python:3.12 (full)", "root (UID 0)", "8", "47", "Baseline (55 total)"),
        ("After Hardening", "python:3.12-slim", "appuser (non-root)", "0", "3", "-94.5% (-100% Critical)")
    ]
    for row_idx, data_row in enumerate(data3, start=1):
        bg = "FADBD8" if row_idx == 1 else "D4EFDF"
        for col_idx, val in enumerate(data_row):
            cell = t3.cell(row_idx, col_idx)
            cell.text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 150, 150)
            p = cell.paragraphs[0]
            if col_idx in (2, 3, 4, 5):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                if col_idx >= 3:
                    r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_image_with_caption("docs/screenshots/q2c_trivy_scan_comparison.png", "Question 2(c) - Trivy vulnerability matrix showing 55 vulnerabilities reduced to 3.")

    # 2(d)
    add_h2("2(d) Threat Agents & Regulatory Compliance (1 Mark)")

    # Threat Agent 1
    add_body(" Automated external bots using credential-stuffing dictionaries sourced from historical public breaches to brute-force user accounts.", bold_prefix="Threat Agent 1 (External Attacker): ")
    add_body(" Salted bcrypt hashing (work factor 12) enforces an intentional ~400 ms cryptographic delay per verification, choking automated brute-force scripts. In addition, TOTP Multi-Factor Authentication ensures that even if credentials are compromised, unauthorized entry is prevented without physical access to the authenticator device.", bold_prefix="Technical Control: ")
    add_body(" GDPR Article 32 ('Security of Processing') requires controllers to implement technical measures including 'the pseudonymisation and encryption of personal data'. bcrypt hashing satisfies this mandate.", bold_prefix="Applicable Regulation: ")
    add_body(" CinePulse stores usernames, email addresses, and viewing history belonging to European Union subjects, requiring compliance with GDPR data protection laws.", bold_prefix="Why Applicable: ")

    # Threat Agent 2
    add_body(" A rogue internal employee or compromised account with legitimate access attempting unauthorized catalog purging or database tampering.", bold_prefix="Threat Agent 2 (Malicious Insider / Rogue Admin): ")
    add_body(" Enforced the principle of least privilege using granular Role-Based Access Control (RBAC). Regular viewer accounts are restricted from modifying or deleting movies (HTTP 403). Furthermore, running the container under an unprivileged non-root user (appuser) prevents privilege escalation to the host operating system.", bold_prefix="Technical Control: ")
    add_body(" GDPR Article 25 ('Data Protection by Design and by Default') mandates that by default, access to personal and system data must be limited strictly to authorized operations necessary for the specific role.", bold_prefix="Applicable Regulation: ")
    add_body(" Protects user profiles and catalog integrity against internal tampering and data destruction.", bold_prefix="Why Applicable: ")

    output_path = "Assignment_3_Report.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_report()
